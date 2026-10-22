#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = [
#   "openpyxl>=3.1.0",
# ]
# ///
"""Normalize a CSV or XLSX Orbit export to patents.json.

读取 xlsx 依赖 openpyxl，已写在本文件顶部。用 `uv run normalize_orbit.py` 执行。csv 不需要它。
"""

from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

ALIASES = {
    "pub_no": ("Publication number", "Publication Number", "公开号", "公开（公告）号"),
    "app_no": ("Application number", "Application Number", "申请号"),
    "title": ("Title", "标题", "发明名称"),
    "applicants": ("Applicant", "Applicants", "Assignee", "申请人"),
    "app_date": ("Application date", "申请日"),
    "pub_date": ("Publication date", "公开日"),
    "grant_date": ("Grant date", "授权日"),
    "legal_status_raw": ("Legal status", "Status", "法律状态"),
    "ipc": ("IPC", "IPC Main", "IPC分类号"),
}


def first(row: dict[str, Any], key: str) -> str:
    for alias in ALIASES[key]:
        value = row.get(alias)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def split_values(value: str) -> list[str]:
    for separator in (";", "|", "\n"):
        value = value.replace(separator, ",")
    return [item.strip() for item in value.split(",") if item.strip()]


def status(value: str) -> str:
    text = value.lower()
    if any(word in text for word in ("pending", "审中", "审查")):
        return "pending"
    if any(word in text for word in ("expired", "过期", "届满", "期满")):
        return "expired"
    if any(word in text for word in ("inactive", "lapsed", "失效", "撤回", "驳回", "无效")):
        return "lapsed"
    return "active"


def patent_type(pub_no: str) -> str:
    upper = pub_no.upper()
    if upper.endswith(("U", "Y")):
        return "utility_model"
    if upper.endswith(("S", "D")):
        return "design"
    return "invention" if upper else "unknown"


def rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 xlsx 需要 openpyxl。请用 uv run 执行本脚本。") from exc
        sheet = load_workbook(path, read_only=True, data_only=True).active
        values = sheet.iter_rows(values_only=True)
        header = [str(value or "").strip() for value in next(values)]
        return [dict(zip(header, row)) for row in values]
    raise RuntimeError("仅支持 .csv 和 .xlsx")


def normalize(row: dict[str, Any]) -> dict[str, Any] | None:
    pub_no = first(row, "pub_no")
    if not pub_no:
        return None
    raw_status = first(row, "legal_status_raw")
    now = datetime.now().astimezone().isoformat(timespec="seconds")
    return {
        "pub_no": pub_no,
        "app_no": first(row, "app_no") or None,
        "family_id": None,
        "title": first(row, "title"),
        "applicants": split_values(first(row, "applicants")),
        "patent_type": patent_type(pub_no),
        "kind_code": pub_no[-1:].upper() or None,
        "patsnap_id": None,
        "jurisdiction": pub_no[:2].upper(),
        "app_date": first(row, "app_date") or None,
        "pub_date": first(row, "pub_date") or None,
        "grant_date": first(row, "grant_date") or None,
        "est_expiry": None,
        "legal_status": status(raw_status),
        "legal_status_raw": {"orbit": raw_status},
        "ipc": split_values(first(row, "ipc")),
        "cpc": [],
        "standardized_assignees": [],
        "claims_available": False,
        "claims_lang": None,
        "source": "orbit_export",
        "found_by": [],
        "fetched_at": now,
        "relevance_rank": None,
        "vendor_summary": {},
        "lifecycle": {},
        "grade": None,
        "grade_reason": "",
        "family_members": [],
        "conflicts": [],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.input.is_file():
        print(f"ORBIT_NORMALIZE: ok=0 missing={args.input}")
        return 2
    try:
        records = [record for row in rows(args.input) if (record := normalize(row))]
    except (OSError, RuntimeError, StopIteration) as exc:
        print(f"ORBIT_NORMALIZE: ok=0 error={exc}")
        return 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"ORBIT_NORMALIZE: ok=1 patents={len(records)} path={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
